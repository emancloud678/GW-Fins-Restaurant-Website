import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"

# ---------------------------------------------------------------------------
# Master data: all 67 original rows, preserved verbatim, recategorized into
# the 10 requested categories, with intent/type/confirmation added.
# ---------------------------------------------------------------------------
# Fields: id, category, intent, question, answer, type, confirm(bool), notes
DATA = [
(1,"Reservations","Make a reservation","Do you take reservations, and how do I book a table?","Yes! We recommend reservations, especially on weekends. You can book directly on our website's Reservations page, or call us at (504) 581-3467.","RAG",False,""),
(2,"Reservations","Book a table via chatbot","Can I book a table right now through this chat?","I can take you straight to our online reservation system — just tell me your preferred date, time, and party size and I'll get you started.","Automation",False,"Connect to live reservation system (OpenTable/Resy/Tock) via n8n."),
(3,"Reservations","Cancellation policy","What is your cancellation policy?","We ask that you cancel or modify at least 24 hours in advance. Large parties (8+) may be subject to a deposit or cancellation fee.","RAG",True,"Confirm exact policy and any deposit amount with restaurant management."),
(4,"Reservations","No-show fee policy","Is there a fee for a no-show?","Please let us know as soon as possible if you can't make it. Repeated no-shows or large-party no-shows may incur a fee — call us to confirm details for your reservation.","RAG",True,"Confirm no-show fee policy with restaurant."),
(5,"Reservations","Request specific table","Can I request a specific table (window, booth, bar)?","We're happy to note table preferences on your reservation, though seating is confirmed on arrival based on availability. I'll add your request to the notes.","Automation",False,""),
(6,"Reservations","Walk-in availability","Do you accept walk-ins?","Yes, walk-ins are welcome and we also offer bar and lounge seating, but a reservation is the best way to guarantee a table, especially in the evening.","RAG",False,""),
(7,"Reservations","Large party booking","What's the largest party size you can accommodate?","For parties of 8 or more, please call us directly at (504) 581-3467 so our events team can arrange the best seating for your group.","RAG",True,"Confirm max walk-in/standard party size threshold."),
(8,"Reservations","Reservation lead time","How far in advance should I book, especially for weekends or holidays?","We recommend booking 1-2 weeks ahead for weekends, and as early as possible for holidays or special events, since we fill up quickly.","RAG",False,""),
(9,"Reservations","Same-day reservation","Can I make a same-day reservation?","Same-day reservations are often available depending on the night — let me check availability for you, or call us directly at (504) 581-3467.","Automation",False,""),
(10,"Reservations","Modify/cancel reservation","How do I modify or cancel an existing reservation?","You can reply here with your reservation details and what you'd like to change, or call us at (504) 581-3467 and we'll update it for you.","Automation",False,""),

(11,"Hours & Location","Operating hours","What are your hours?","We're open for dinner nightly starting at 5:00 PM. We do not receive our daily catch until 4:00 PM, so our full menu and service begin at 5 PM.","RAG",False,""),
(12,"Hours & Location","Lunch availability","Are you open for lunch?","We currently serve dinner only, nightly from 5:00 PM.","RAG",True,"Confirm — some listings reference lunch/takeout; verify current hours with restaurant."),
(13,"Hours & Location","Restaurant location","Where are you located?","We're located at 808 Bienville St, New Orleans, LA 70112 — in the historic DH Holmes building, just 4 blocks from the Saenger Theatre.","RAG",False,""),
(14,"Hours & Location","Holiday hours","Are you open on holidays?","Our holiday hours can vary — please call us at (504) 581-3467 or check our website closer to the date for confirmed holiday hours.","RAG",True,"Client to supply holiday hours calendar."),
(15,"Hours & Location","Pre-theater dining","Is GW Fins good for a pre-theater dinner?","Absolutely — we're only 4 blocks from the Saenger Theatre, making us a favorite stop for dinner before a show.","RAG",False,""),
(16,"Hours & Location","Neighborhood/area","What neighborhood are you in?","We're in the Central Business District, right at the edge of the French Quarter, inside the historic DH Holmes building on Bienville Street.","RAG",False,""),
(17,"Hours & Location","Contact restaurant","How do I contact the restaurant directly?","You can call us at (504) 581-3467, or I can pass your message along — just let me know what you need.","Automation",False,""),
(18,"Hours & Location","Website info","Do you have a website I can browse?","Yes, you can see our full menu, wine list, and book a table at gwfins.com.","RAG",False,""),

(19,"Menu & Dining","Cuisine type","What kind of food do you serve?","We're a premier seafood restaurant featuring the freshest Gulf and global catch, prepared with modern technique alongside New Orleans classics like our Scalibut and Barbeque Shrimp.","RAG",False,""),
(20,"Menu & Dining","Menu changes daily","Does your menu change?","Yes — our menu changes daily since we don't receive that day's fresh catch until 4:00 PM, so you'll always get the freshest fish available.","RAG",False,""),
(21,"Menu & Dining","View today's menu","Can I see today's menu?","I can share our most recently posted menu, but please note the fresh fish selections may vary slightly by the day. You can also view it live at gwfins.com/menu.","Automation",False,""),
(22,"Menu & Dining","Vegetarian/vegan options","Do you have vegetarian or vegan options?","We have salads and side dishes that can be made vegetarian, but our menu is seafood-focused. Let your server know of any dietary needs and our kitchen will do its best to accommodate.","RAG",True,"Confirm dedicated vegetarian/vegan dish availability with chef."),
(23,"Menu & Dining","Food allergies","Do you accommodate food allergies (shellfish, gluten, nuts)?","Yes, please inform your server of any allergies when you arrive, and our kitchen team will guide you to safe options. Given we're a seafood restaurant, please flag shellfish allergies clearly.","RAG",True,"Confirm formal allergy protocol with restaurant for liability accuracy."),
(24,"Menu & Dining","Scalibut dish info","What is 'Scalibut'?","Scalibut is a GW Fins original: halibut and sea scallops served over Royal Red shrimp risotto with snow peas and pea shoot butter — one of our most popular dishes.","RAG",False,""),
(25,"Menu & Dining","Kids menu","Do you offer a kids' menu?","We don't have a dedicated printed kids' menu, but our kitchen is happy to prepare a simple dish for younger guests — just ask your server.","RAG",True,"Confirm with restaurant whether a kids' menu exists."),
(26,"Menu & Dining","Popular dishes","What's your most popular dish?","Guests especially love our Scalibut, Lobster Dumplings, and Lobster Bisque — all fan favorites!","RAG",False,""),
(27,"Menu & Dining","Gluten-free options","Do you offer gluten-free options?","Many of our seafood entrées can be prepared gluten-friendly — let your server know and they'll point you to the best options.","RAG",True,"Confirm formal gluten-free protocol with restaurant."),
(28,"Menu & Dining","Wine/cocktail list","Can I see the wine and cocktail list?","Of course — we have a full wine list (updated daily) and a specialty cocktail menu featuring drinks like our French 75 and Whistle Old Fashioned. You can view both at gwfins.com/menu.","RAG",False,""),

(29,"Pricing & Payment","Dinner pricing","How much does dinner cost?","Entrées typically range from about $32-$56, with most guests spending around $40-$100 per person depending on courses and drinks.","RAG",False,""),
(30,"Pricing & Payment","Payment methods","Do you accept credit cards?","Yes, we accept all major credit cards.","RAG",True,"Confirm accepted card types/payment methods with restaurant."),
(31,"Pricing & Payment","Gift card purchase","Do you offer gift cards?","Yes! Gift cards are available — I can point you to where to purchase them on our website, or you can call us directly.","RAG",True,"Confirm gift card purchase link/process with restaurant."),
(32,"Pricing & Payment","Corkage fee","Is there a corkage fee if I bring my own wine?","Please call us at (504) 581-3467 to confirm our current corkage policy and fee.","RAG",True,"Confirm corkage policy with restaurant."),
(33,"Pricing & Payment","Automatic gratuity","Do you add automatic gratuity?","Automatic gratuity may apply for larger parties (typically 6-8+). Please confirm with your server or call ahead for group bookings.","RAG",True,"Confirm gratuity threshold/percentage with restaurant."),
(34,"Pricing & Payment","Tax/tip inclusion","Is tax and tip included in menu prices?","Menu prices do not include tax or gratuity.","RAG",False,""),
(56,"Pricing & Payment","Buy gift card","How do I buy a gift card?","You can purchase a gift card through our website or by calling us directly at (504) 581-3467.","RAG",True,"Confirm gift card purchase link with restaurant."),
(57,"Pricing & Payment","Gift card expiration","Do gift cards expire?","Please call us at (504) 581-3467 to confirm current gift card terms.","RAG",True,"Confirm expiration policy with restaurant."),
(58,"Pricing & Payment","Loyalty program","Do you have a rewards or loyalty program?","Please call us at (504) 581-3467 or ask your server about any current loyalty or repeat-guest perks.","RAG",True,"Confirm if a loyalty program exists."),

(35,"Private Events & Groups","Private event hosting","Do you host private events or parties?","Yes, we offer private dining spaces perfect for celebrations, corporate events, and special occasions. Visit our Private Parties page or call us to check availability.","RAG",False,""),
(36,"Private Events & Groups","Private dining capacity","What is the capacity of your private dining room?","Our private dining spaces can accommodate a range of group sizes — call us at (504) 581-3467 and our events team will help find the right fit.","RAG",True,"Confirm exact private room capacities with restaurant."),
(37,"Private Events & Groups","Birthday celebration","Can I bring a birthday cake or celebrate a special occasion?","Absolutely — let us know when booking that you're celebrating, and we'll do our best to make it special. Outside desserts are usually fine with advance notice.","RAG",True,"Confirm outside cake/dessert policy with restaurant."),
(38,"Private Events & Groups","Corporate/rehearsal dinner planning","How do I plan a rehearsal dinner or corporate event?","We'd love to help! Please call (504) 581-3467 or reach out through our Private Parties page and our events coordinator will assist with menus and setup.","Automation",False,""),
(39,"Private Events & Groups","Private event minimum spend","Is there a minimum spend for private events?","Minimum spend depends on the date, room, and group size — our events team can give you exact details when you inquire.","RAG",True,"Confirm minimum spend policy with restaurant."),
(40,"Private Events & Groups","Large group seating","Can you accommodate a large group without a private room?","Yes, for groups of 8+ we can often arrange seating in our main dining room — call ahead so we can prepare the best table for you.","RAG",False,""),

(41,"Dress Code & Atmosphere","Dress code","What is the dress code?","We recommend business casual to upscale attire — collared shirts for gentlemen are appreciated, though we don't strictly require formal wear.","RAG",True,"Confirm exact dress code policy with restaurant."),
(42,"Dress Code & Atmosphere","Formality/atmosphere","Is GW Fins formal or casual?","We're upscale but approachable — an elegant dining room with attentive service, great for both special occasions and a nice night out.","RAG",False,""),
(43,"Dress Code & Atmosphere","Ambiance (loud/quiet)","Is it loud/lively or quiet and romantic?","Guests describe our atmosphere as elegant and lively — a great choice for both celebrations and more intimate dinners.","RAG",False,""),
(44,"Dress Code & Atmosphere","Outdoor/patio seating","Do you have outdoor or patio seating?","Please call us at (504) 581-3467 to confirm current seating options, including any patio or bar seating availability.","RAG",True,"Confirm outdoor/patio seating availability with restaurant."),

(45,"Parking & Accessibility","Parking availability","Is there parking available?","There is parking available near the restaurant in the DH Holmes building and surrounding CBD lots/garages. Please call us at (504) 581-3467 for current valet or parking recommendations.","RAG",True,"Confirm valet availability/validated parking with restaurant."),
(46,"Parking & Accessibility","Valet parking","Do you offer valet parking?","Please call us at (504) 581-3467 to confirm valet availability, as this can vary by night.","RAG",True,"Confirm valet policy with restaurant."),
(47,"Parking & Accessibility","Wheelchair accessibility","Is the restaurant wheelchair accessible?","Yes, please let us know if you need any accessibility accommodations when booking, and we'll make sure you're well taken care of.","RAG",True,"Confirm ADA accessibility details with restaurant."),
(48,"Parking & Accessibility","Proximity to hotels","Are you close to French Quarter hotels?","Yes, we're right at the edge of the French Quarter in the CBD, an easy walk from most major hotels and just 4 blocks from the Saenger Theatre.","RAG",False,""),
(49,"Parking & Accessibility","Pet policy","Do you allow pets/service animals?","Service animals are always welcome. For general pets, please call ahead to confirm any patio pet policy.","RAG",True,"Confirm pet policy with restaurant."),
(50,"Parking & Accessibility","Walkability from Bourbon St","Is there a shuttle or is it walkable from Bourbon Street?","We're very walkable from Bourbon Street and the French Quarter — just a few minutes on foot.","RAG",False,""),

(51,"Bar & Wine","Bar-only visit","Can I just come for drinks without dinner?","Yes! Our bar and lounge area is open for cocktails, wine, and light bites even if you're not staying for a full dinner.","RAG",True,"Confirm bar-only seating policy with restaurant."),
(52,"Bar & Wine","Cocktail menu","What's on your cocktail menu?","We feature specialty cocktails like the French 75, A Spicy Marg?, Peach Hugo Spritz, and the Whistle Old Fashioned, plus zero-proof options like our Not At All Spritz.","RAG",False,""),
(53,"Bar & Wine","Non-alcoholic drinks","Do you have non-alcoholic drink options?","Yes, we offer zero-proof cocktails like the Not At All Spritz and Watermelon No-Jito, along with French Truck coffee service.","RAG",False,""),
(54,"Bar & Wine","Wine list scope","Do you have an extensive wine list?","Yes, our wine list changes daily and features 100+ selections across sparkling, white, rosé, and red. Ask your server for today's list or view featured by-the-glass pours on our site.","RAG",False,""),
(55,"Bar & Wine","Happy hour","Do you have happy hour?","Please call us at (504) 581-3467 to confirm if a happy hour or bar special is currently running.","RAG",True,"Confirm happy hour program with restaurant."),

(59,"Special Occasions","Anniversary/proposal dining","Is GW Fins good for anniversaries or proposals?","Yes, we're a popular choice for anniversaries, birthdays, and proposals — let us know when booking and we'll help make it memorable.","RAG",False,""),
(60,"Special Occasions","Celebration extras","Can you arrange a special dessert or celebration touch?","Yes, just mention the occasion when booking or let your server know, and we'll do our best to make it special.","RAG",True,"Confirm what complimentary celebration touches are offered."),
(61,"Special Occasions","Restaurant history/anniversary","Do you celebrate a restaurant anniversary or special history?","Yes! GW Fins is celebrating its 25th Anniversary (2001-2026) as one of New Orleans' premier fine dining seafood restaurants.","RAG",False,""),
(62,"Special Occasions","Business dinner suitability","Is GW Fins good for business dinners?","Yes, our elegant dining room and attentive service make it a great choice for business dinners, and we also offer private event spaces for larger meetings.","RAG",False,""),

(63,"Contact & General","Speak to human staff","How do I speak to a real person?","Of course — you can call us directly at (504) 581-3467, and our team will be happy to help with anything I can't answer here.","Human Handoff",False,""),
(64,"Contact & General","Phone number","What's your phone number?","You can reach us at (504) 581-3467.","RAG",False,""),
(65,"Contact & General","Social media","Do you have social media I can follow?","Yes, you can find us on Instagram and Facebook @gwfinsneworleans for daily menu highlights and updates.","RAG",True,"Confirm exact social handles with restaurant."),
(66,"Contact & General","Leave feedback/review","Can I leave feedback or a review?","We'd love to hear from you! You can leave a review on Google or reach out directly — we read every comment.","RAG",False,""),
(67,"Contact & General","Takeout/delivery availability","Do you offer takeout or delivery?","Please call us at (504) 581-3467 to confirm current takeout availability, as GW Fins is primarily a dine-in fine dining experience.","RAG",True,"Confirm takeout/delivery availability — listings vary."),
]

assert len(DATA) == 67, f"expected 67 rows, got {len(DATA)}"
assert sorted(d[0] for d in DATA) == list(range(1,68)), "IDs 1-67 must all be present exactly once"

CATEGORY_ORDER = ["Reservations","Hours & Location","Menu & Dining","Pricing & Payment",
                   "Private Events & Groups","Dress Code & Atmosphere","Parking & Accessibility",
                   "Bar & Wine","Special Occasions","Contact & General"]
for d in DATA:
    assert d[1] in CATEGORY_ORDER, f"bad category on row {d[0]}: {d[1]}"

AUTOMATION_MAP = [
(2,"Book table via chatbot","Can I book a table right now through this chat?",
 "Capture date, time, party size, and contact info; create a live reservation.",
 "Chat trigger -> collect booking fields -> HTTP Request node to reservation platform API (OpenTable/Resy/Tock) -> confirm booking back to guest -> optional confirmation email/SMS.",
 "Confirm which reservation platform/API the restaurant uses."),
(5,"Request specific table","Can I request a specific table (window, booth, bar)?",
 "Attach a free-text seating preference note to the reservation record.",
 "Extend the booking workflow (ID 2) to pass a 'special_requests' field into the reservation API call or booking notes field.",
 "Depends on reservation platform supporting a notes/special-requests field."),
(9,"Same-day reservation","Can I make a same-day reservation?",
 "Check live table availability for today before answering.",
 "HTTP Request node -> reservation platform availability endpoint filtered to today's date -> return open time slots to guest.",
 "Requires reservation platform API access with a real-time availability endpoint."),
(10,"Modify/cancel reservation","How do I modify or cancel an existing reservation?",
 "Look up an existing reservation and update or cancel it.",
 "Webhook -> look up reservation via API (search by phone/email) -> if found, call update/cancel endpoint -> confirm change to guest; fall back to Human Handoff if not found or platform lacks self-service API.",
 "Consider an identity verification step (phone/email match) before allowing changes."),
(17,"Contact restaurant","How do I contact the restaurant directly?",
 "Route a guest's message to restaurant staff when they want more than the phone number.",
 "If guest opts to leave a message -> capture message + contact info -> send via Email/Slack/SMS node to restaurant staff inbox.",
 "Confirm which staff inbox/channel should receive these messages."),
(21,"View today's menu","Can I see today's menu?",
 "Serve the current day's menu rather than a static cached answer, since the real menu changes daily after 4 PM.",
 "Scheduled workflow (e.g. daily ~4:30 PM) fetches gwfins.com/menu -> updates the 'current menu' data source / RAG index the chatbot queries.",
 "Needs a reliable feed or scrape source; confirm with client if an API/export is available instead of scraping."),
(38,"Corporate/rehearsal dinner planning","How do I plan a rehearsal dinner or corporate event?",
 "Capture event details (date, group size, event type, contact info) and forward as a lead to the events coordinator.",
 "Conversational form collects event details -> n8n sends lead via Email/CRM node to events coordinator -> auto-reply confirmation to guest.",
 "Confirm preferred lead destination (email address or CRM) with restaurant."),
(63,"Speak to human staff","How do I speak to a real person?",
 "Escalate the conversation to restaurant staff instead of continuing with bot answers.",
 "Human Handoff branch: notify staff via Slack/SMS/email with the chat transcript, or surface the phone number prominently with a one-tap call action.",
 "Decide whether true live handoff (staff joins chat) is in scope, or the bot simply always surfaces the phone number."),
]

# ---------------------------------------------------------------------------
# Styling helpers
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
CATEGORY_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
CONFIRM_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
TYPE_FILLS = {
    "RAG": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
    "Automation": PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
    "Human Handoff": PatternFill(start_color="E1D5E7", end_color="E1D5E7", fill_type="solid"),
}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def style_header(ws, ncols, row=1, height=30):
    for col in range(1, ncols+1):
        c = ws.cell(row=row, column=col)
        c.font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = height
    ws.freeze_panes = ws.cell(row=row+1, column=1).coordinate

wb = openpyxl.Workbook()

# ---------- Sheet 1: MASTER FAQ ----------
ws1 = wb.active
ws1.title = "MASTER FAQ"
headers1 = ["ID","Category","Customer Intent","Customer Question","Suggested Chatbot Answer",
            "Type (RAG / Automation / Human Handoff)","Client Confirmation Required","Notes"]
ws1.append(headers1)
style_header(ws1, len(headers1))

# Order rows by the requested category order, preserving original IDs
ordered = sorted(DATA, key=lambda d: (CATEGORY_ORDER.index(d[1]), d[0]))

r = 2
for (rid, cat, intent, q, a, typ, confirm, note) in ordered:
    ws1.cell(row=r, column=1, value=rid)
    cat_cell = ws1.cell(row=r, column=2, value=cat)
    cat_cell.fill = CATEGORY_FILL
    ws1.cell(row=r, column=3, value=intent)
    q_cell = ws1.cell(row=r, column=4, value=q)
    q_cell.font = Font(name=FONT_NAME, size=10, bold=True)
    ws1.cell(row=r, column=5, value=a)
    type_cell = ws1.cell(row=r, column=6, value=typ)
    type_cell.fill = TYPE_FILLS[typ]
    type_cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    confirm_cell = ws1.cell(row=r, column=7, value="Yes" if confirm else "No")
    confirm_cell.alignment = Alignment(horizontal="center", vertical="top")
    if confirm:
        confirm_cell.fill = CONFIRM_FILL
    note_cell = ws1.cell(row=r, column=8, value=note)
    note_cell.font = Font(name=FONT_NAME, size=9, italic=True, color="7F6000")

    for col in range(1, len(headers1)+1):
        c = ws1.cell(row=r, column=col)
        if c.font.name is None:
            c.font = Font(name=FONT_NAME, size=10)
        c.border = BORDER
        if col not in (6,7):
            c.alignment = Alignment(vertical="top", wrap_text=True)
    r += 1

ws1.column_dimensions["A"].width = 6
ws1.column_dimensions["B"].width = 20
ws1.column_dimensions["C"].width = 24
ws1.column_dimensions["D"].width = 38
ws1.column_dimensions["E"].width = 48
ws1.column_dimensions["F"].width = 16
ws1.column_dimensions["G"].width = 14
ws1.column_dimensions["H"].width = 34
ws1.auto_filter.ref = f"A1:H{r-1}"

# ---------- Sheet 2: CLIENT CONFIRMATION ----------
ws2 = wb.create_sheet("CLIENT CONFIRMATION")
headers2 = ["ID","Category","Question","Information That Needs Confirmation","Status"]
ws2.append(headers2)
style_header(ws2, len(headers2))

confirm_rows = [d for d in DATA if d[6]]
confirm_rows.sort(key=lambda d: (CATEGORY_ORDER.index(d[1]), d[0]))

r = 2
for (rid, cat, intent, q, a, typ, confirm, note) in confirm_rows:
    ws2.cell(row=r, column=1, value=rid)
    cat_cell = ws2.cell(row=r, column=2, value=cat)
    cat_cell.fill = CATEGORY_FILL
    ws2.cell(row=r, column=3, value=q)
    ws2.cell(row=r, column=4, value=note)
    status_cell = ws2.cell(row=r, column=5, value="Needs Client Confirmation")
    status_cell.fill = CONFIRM_FILL
    status_cell.font = Font(name=FONT_NAME, size=10, bold=True, color="7F6000")
    status_cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    for col in range(1, len(headers2)+1):
        c = ws2.cell(row=r, column=col)
        if c.font is None or c.font.name is None:
            c.font = Font(name=FONT_NAME, size=10)
        c.border = BORDER
        if col != 5:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    r += 1

ws2.column_dimensions["A"].width = 6
ws2.column_dimensions["B"].width = 20
ws2.column_dimensions["C"].width = 42
ws2.column_dimensions["D"].width = 46
ws2.column_dimensions["E"].width = 22
ws2.auto_filter.ref = f"A1:E{r-1}"

# ---------- Sheet 3: AUTOMATION MAP ----------
ws3 = wb.create_sheet("AUTOMATION MAP")
headers3 = ["ID","Intent","Customer Request","Required Action","Suggested n8n Automation","Notes"]
ws3.append(headers3)
style_header(ws3, len(headers3))

r = 2
for (rid, intent, req, action, automation, note) in AUTOMATION_MAP:
    ws3.cell(row=r, column=1, value=rid)
    ws3.cell(row=r, column=2, value=intent).font = Font(name=FONT_NAME, size=10, bold=True)
    ws3.cell(row=r, column=3, value=req)
    ws3.cell(row=r, column=4, value=action)
    ws3.cell(row=r, column=5, value=automation)
    note_cell = ws3.cell(row=r, column=6, value=note)
    note_cell.font = Font(name=FONT_NAME, size=9, italic=True, color="7F6000")
    for col in range(1, len(headers3)+1):
        c = ws3.cell(row=r, column=col)
        if c.font is None or c.font.name is None:
            c.font = Font(name=FONT_NAME, size=10)
        c.border = BORDER
        c.alignment = Alignment(vertical="top", wrap_text=True)
    r += 1

ws3.column_dimensions["A"].width = 6
ws3.column_dimensions["B"].width = 24
ws3.column_dimensions["C"].width = 36
ws3.column_dimensions["D"].width = 38
ws3.column_dimensions["E"].width = 48
ws3.column_dimensions["F"].width = 34
ws3.auto_filter.ref = f"A1:F{r-1}"

out_path = r"C:\Users\user\Downloads\GW Fins restaurant\chatbot\GW_Fins_Chatbot_FAQ_v2.xlsx"
wb.save(out_path)
print("saved", out_path)
print("Master rows:", len(ordered), "| Confirmation rows:", len(confirm_rows), "| Automation rows:", len(AUTOMATION_MAP))
