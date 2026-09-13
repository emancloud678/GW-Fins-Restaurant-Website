import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ---------- Sheet 1: FAQ Knowledge Base ----------
ws = wb.active
ws.title = "FAQ Knowledge Base"

FONT_NAME = "Arial"
HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
CATEGORY_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
NOTE_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

headers = ["#", "Category", "Customer Question (Intent)", "Suggested Chatbot Answer", "Notes / Needs Client Confirmation"]
ws.append(headers)
for col, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=col)
    c.font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
    c.fill = HEADER_FILL
    c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    c.border = BORDER
ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 28

# Each row: category, question, answer, note (note may be "" )
rows = [
("Reservations", "Do you take reservations, and how do I book a table?",
 "Yes! We recommend reservations, especially on weekends. You can book directly on our website's Reservations page, or call us at (504) 581-3467.",
 ""),
("Reservations", "Can I book a table right now through this chat?",
 "I can take you straight to our online reservation system — just tell me your preferred date, time, and party size and I'll get you started.",
 "Connect to live reservation system (OpenTable/Resy/Tock) via n8n."),
("Reservations", "What is your cancellation policy?",
 "We ask that you cancel or modify at least 24 hours in advance. Large parties (8+) may be subject to a deposit or cancellation fee.",
 "Confirm exact policy and any deposit amount with restaurant management."),
("Reservations", "Is there a fee for a no-show?",
 "Please let us know as soon as possible if you can't make it. Repeated no-shows or large-party no-shows may incur a fee — call us to confirm details for your reservation.",
 "Confirm no-show fee policy with restaurant."),
("Reservations", "Can I request a specific table (window, booth, bar)?",
 "We're happy to note table preferences on your reservation, though seating is confirmed on arrival based on availability. I'll add your request to the notes.",
 ""),
("Reservations", "Do you accept walk-ins?",
 "Yes, walk-ins are welcome and we also offer bar and lounge seating, but a reservation is the best way to guarantee a table, especially in the evening.",
 ""),
("Reservations", "What's the largest party size you can accommodate?",
 "For parties of 8 or more, please call us directly at (504) 581-3467 so our events team can arrange the best seating for your group.",
 "Confirm max walk-in/standard party size threshold."),
("Reservations", "How far in advance should I book, especially for weekends or holidays?",
 "We recommend booking 1-2 weeks ahead for weekends, and as early as possible for holidays or special events, since we fill up quickly.",
 ""),
("Reservations", "Can I make a same-day reservation?",
 "Same-day reservations are often available depending on the night — let me check availability for you, or call us directly at (504) 581-3467.",
 ""),
("Reservations", "How do I modify or cancel an existing reservation?",
 "You can reply here with your reservation details and what you'd like to change, or call us at (504) 581-3467 and we'll update it for you.",
 ""),

("Hours & Location", "What are your hours?",
 "We're open for dinner nightly starting at 5:00 PM. We do not receive our daily catch until 4:00 PM, so our full menu and service begin at 5 PM.",
 ""),
("Hours & Location", "Are you open for lunch?",
 "We currently serve dinner only, nightly from 5:00 PM.",
 "Confirm — some listings reference lunch/takeout; verify current hours with restaurant."),
("Hours & Location", "Where are you located?",
 "We're located at 808 Bienville St, New Orleans, LA 70112 — in the historic DH Holmes building, just 4 blocks from the Saenger Theatre.",
 ""),
("Hours & Location", "Are you open on holidays?",
 "Our holiday hours can vary — please call us at (504) 581-3467 or check our website closer to the date for confirmed holiday hours.",
 "Client to supply holiday hours calendar."),
("Hours & Location", "Is GW Fins good for a pre-theater dinner?",
 "Absolutely — we're only 4 blocks from the Saenger Theatre, making us a favorite stop for dinner before a show.",
 ""),
("Hours & Location", "What neighborhood are you in?",
 "We're in the Central Business District, right at the edge of the French Quarter, inside the historic DH Holmes building on Bienville Street.",
 ""),
("Hours & Location", "How do I contact the restaurant directly?",
 "You can call us at (504) 581-3467, or I can pass your message along — just let me know what you need.",
 ""),
("Hours & Location", "Do you have a website I can browse?",
 "Yes, you can see our full menu, wine list, and book a table at gwfins.com.",
 ""),

("Menu & Dietary", "What kind of food do you serve?",
 "We're a premier seafood restaurant featuring the freshest Gulf and global catch, prepared with modern technique alongside New Orleans classics like our Scalibut and Barbeque Shrimp.",
 ""),
("Menu & Dietary", "Does your menu change?",
 "Yes — our menu changes daily since we don't receive that day's fresh catch until 4:00 PM, so you'll always get the freshest fish available.",
 ""),
("Menu & Dietary", "Can I see today's menu?",
 "I can share our most recently posted menu, but please note the fresh fish selections may vary slightly by the day. You can also view it live at gwfins.com/menu.",
 ""),
("Menu & Dietary", "Do you have vegetarian or vegan options?",
 "We have salads and side dishes that can be made vegetarian, but our menu is seafood-focused. Let your server know of any dietary needs and our kitchen will do its best to accommodate.",
 "Confirm dedicated vegetarian/vegan dish availability with chef."),
("Menu & Dietary", "Do you accommodate food allergies (shellfish, gluten, nuts)?",
 "Yes, please inform your server of any allergies when you arrive, and our kitchen team will guide you to safe options. Given we're a seafood restaurant, please flag shellfish allergies clearly.",
 "Confirm formal allergy protocol with restaurant for liability accuracy."),
("Menu & Dietary", "What is 'Scalibut'?",
 "Scalibut is a GW Fins original: halibut and sea scallops served over Royal Red shrimp risotto with snow peas and pea shoot butter — one of our most popular dishes.",
 ""),
("Menu & Dietary", "Do you offer a kids' menu?",
 "We don't have a dedicated printed kids' menu, but our kitchen is happy to prepare a simple dish for younger guests — just ask your server.",
 "Confirm with restaurant whether a kids' menu exists."),
("Menu & Dietary", "What's your most popular dish?",
 "Guests especially love our Scalibut, Lobster Dumplings, and Lobster Bisque — all fan favorites!",
 ""),
("Menu & Dietary", "Do you offer gluten-free options?",
 "Many of our seafood entrées can be prepared gluten-friendly — let your server know and they'll point you to the best options.",
 "Confirm formal gluten-free protocol with restaurant."),
("Menu & Dietary", "Can I see the wine and cocktail list?",
 "Of course — we have a full wine list (updated daily) and a specialty cocktail menu featuring drinks like our French 75 and Whistle Old Fashioned. You can view both at gwfins.com/menu.",
 ""),

("Pricing & Payment", "How much does dinner cost?",
 "Entrées typically range from about $32-$56, with most guests spending around $40-$100 per person depending on courses and drinks.",
 ""),
("Pricing & Payment", "Do you accept credit cards?",
 "Yes, we accept all major credit cards.",
 "Confirm accepted card types/payment methods with restaurant."),
("Pricing & Payment", "Do you offer gift cards?",
 "Yes! Gift cards are available — I can point you to where to purchase them on our website, or you can call us directly.",
 "Confirm gift card purchase link/process with restaurant."),
("Pricing & Payment", "Is there a corkage fee if I bring my own wine?",
 "Please call us at (504) 581-3467 to confirm our current corkage policy and fee.",
 "Confirm corkage policy with restaurant."),
("Pricing & Payment", "Do you add automatic gratuity?",
 "Automatic gratuity may apply for larger parties (typically 6-8+). Please confirm with your server or call ahead for group bookings.",
 "Confirm gratuity threshold/percentage with restaurant."),
("Pricing & Payment", "Is tax and tip included in menu prices?",
 "Menu prices do not include tax or gratuity.",
 ""),

("Private Events & Groups", "Do you host private events or parties?",
 "Yes, we offer private dining spaces perfect for celebrations, corporate events, and special occasions. Visit our Private Parties page or call us to check availability.",
 ""),
("Private Events & Groups", "What is the capacity of your private dining room?",
 "Our private dining spaces can accommodate a range of group sizes — call us at (504) 581-3467 and our events team will help find the right fit.",
 "Confirm exact private room capacities with restaurant."),
("Private Events & Groups", "Can I bring a birthday cake or celebrate a special occasion?",
 "Absolutely — let us know when booking that you're celebrating, and we'll do our best to make it special. Outside desserts are usually fine with advance notice.",
 "Confirm outside cake/dessert policy with restaurant."),
("Private Events & Groups", "How do I plan a rehearsal dinner or corporate event?",
 "We'd love to help! Please call (504) 581-3467 or reach out through our Private Parties page and our events coordinator will assist with menus and setup.",
 ""),
("Private Events & Groups", "Is there a minimum spend for private events?",
 "Minimum spend depends on the date, room, and group size — our events team can give you exact details when you inquire.",
 "Confirm minimum spend policy with restaurant."),
("Private Events & Groups", "Can you accommodate a large group without a private room?",
 "Yes, for groups of 8+ we can often arrange seating in our main dining room — call ahead so we can prepare the best table for you.",
 ""),

("Dress Code & Atmosphere", "What is the dress code?",
 "We recommend business casual to upscale attire — collared shirts for gentlemen are appreciated, though we don't strictly require formal wear.",
 "Confirm exact dress code policy with restaurant."),
("Dress Code & Atmosphere", "Is GW Fins formal or casual?",
 "We're upscale but approachable — an elegant dining room with attentive service, great for both special occasions and a nice night out.",
 ""),
("Dress Code & Atmosphere", "Is it loud/lively or quiet and romantic?",
 "Guests describe our atmosphere as elegant and lively — a great choice for both celebrations and more intimate dinners.",
 ""),
("Dress Code & Atmosphere", "Do you have outdoor or patio seating?",
 "Please call us at (504) 581-3467 to confirm current seating options, including any patio or bar seating availability.",
 "Confirm outdoor/patio seating availability with restaurant."),

("Parking & Accessibility", "Is there parking available?",
 "There is parking available near the restaurant in the DH Holmes building and surrounding CBD lots/garages. Please call us at (504) 581-3467 for current valet or parking recommendations.",
 "Confirm valet availability/validated parking with restaurant."),
("Parking & Accessibility", "Do you offer valet parking?",
 "Please call us at (504) 581-3467 to confirm valet availability, as this can vary by night.",
 "Confirm valet policy with restaurant."),
("Parking & Accessibility", "Is the restaurant wheelchair accessible?",
 "Yes, please let us know if you need any accessibility accommodations when booking, and we'll make sure you're well taken care of.",
 "Confirm ADA accessibility details with restaurant."),
("Parking & Accessibility", "Are you close to French Quarter hotels?",
 "Yes, we're right at the edge of the French Quarter in the CBD, an easy walk from most major hotels and just 4 blocks from the Saenger Theatre.",
 ""),
("Parking & Accessibility", "Do you allow pets/service animals?",
 "Service animals are always welcome. For general pets, please call ahead to confirm any patio pet policy.",
 "Confirm pet policy with restaurant."),
("Parking & Accessibility", "Is there a shuttle or is it walkable from Bourbon Street?",
 "We're very walkable from Bourbon Street and the French Quarter — just a few minutes on foot.",
 ""),

("Bar / Wine / Cocktails", "Can I just come for drinks without dinner?",
 "Yes! Our bar and lounge area is open for cocktails, wine, and light bites even if you're not staying for a full dinner.",
 "Confirm bar-only seating policy with restaurant."),
("Bar / Wine / Cocktails", "What's on your cocktail menu?",
 "We feature specialty cocktails like the French 75, A Spicy Marg?, Peach Hugo Spritz, and the Whistle Old Fashioned, plus zero-proof options like our Not At All Spritz.",
 ""),
("Bar / Wine / Cocktails", "Do you have non-alcoholic drink options?",
 "Yes, we offer zero-proof cocktails like the Not At All Spritz and Watermelon No-Jito, along with French Truck coffee service.",
 ""),
("Bar / Wine / Cocktails", "Do you have an extensive wine list?",
 "Yes, our wine list changes daily and features 100+ selections across sparkling, white, rosé, and red. Ask your server for today's list or view featured by-the-glass pours on our site.",
 ""),
("Bar / Wine / Cocktails", "Do you have happy hour?",
 "Please call us at (504) 581-3467 to confirm if a happy hour or bar special is currently running.",
 "Confirm happy hour program with restaurant."),

("Gift Cards & Loyalty", "How do I buy a gift card?",
 "You can purchase a gift card through our website or by calling us directly at (504) 581-3467.",
 "Confirm gift card purchase link with restaurant."),
("Gift Cards & Loyalty", "Do gift cards expire?",
 "Please call us at (504) 581-3467 to confirm current gift card terms.",
 "Confirm expiration policy with restaurant."),
("Gift Cards & Loyalty", "Do you have a rewards or loyalty program?",
 "Please call us at (504) 581-3467 or ask your server about any current loyalty or repeat-guest perks.",
 "Confirm if a loyalty program exists."),

("Special Occasions", "Is GW Fins good for anniversaries or proposals?",
 "Yes, we're a popular choice for anniversaries, birthdays, and proposals — let us know when booking and we'll help make it memorable.",
 ""),
("Special Occasions", "Can you arrange a special dessert or celebration touch?",
 "Yes, just mention the occasion when booking or let your server know, and we'll do our best to make it special.",
 "Confirm what complimentary celebration touches are offered."),
("Special Occasions", "Do you celebrate a restaurant anniversary or special history?",
 "Yes! GW Fins is celebrating its 25th Anniversary (2001-2026) as one of New Orleans' premier fine dining seafood restaurants.",
 ""),
("Special Occasions", "Is GW Fins good for business dinners?",
 "Yes, our elegant dining room and attentive service make it a great choice for business dinners, and we also offer private event spaces for larger meetings.",
 ""),

("Contact & General", "How do I speak to a real person?",
 "Of course — you can call us directly at (504) 581-3467, and our team will be happy to help with anything I can't answer here.",
 ""),
("Contact & General", "What's your phone number?",
 "You can reach us at (504) 581-3467.",
 ""),
("Contact & General", "Do you have social media I can follow?",
 "Yes, you can find us on Instagram and Facebook @gwfinsneworleans for daily menu highlights and updates.",
 "Confirm exact social handles with restaurant."),
("Contact & General", "Can I leave feedback or a review?",
 "We'd love to hear from you! You can leave a review on Google or reach out directly — we read every comment.",
 ""),
("Contact & General", "Do you offer takeout or delivery?",
 "Please call us at (504) 581-3467 to confirm current takeout availability, as GW Fins is primarily a dine-in fine dining experience.",
 "Confirm takeout/delivery availability — listings vary."),
]

row_idx = 2
current_cat = None
n = 1
for cat, q, a, note in rows:
    ws.cell(row=row_idx, column=1, value=n).font = Font(name=FONT_NAME, size=10)
    cat_cell = ws.cell(row=row_idx, column=2, value=cat)
    cat_cell.font = Font(name=FONT_NAME, size=10, bold=True)
    cat_cell.fill = CATEGORY_FILL
    q_cell = ws.cell(row=row_idx, column=3, value=q)
    q_cell.font = Font(name=FONT_NAME, size=10, bold=True)
    a_cell = ws.cell(row=row_idx, column=4, value=a)
    a_cell.font = Font(name=FONT_NAME, size=10)
    note_cell = ws.cell(row=row_idx, column=5, value=note)
    note_cell.font = Font(name=FONT_NAME, size=9, italic=True, color="7F6000")
    if note:
        note_cell.fill = NOTE_FILL
    for col in range(1, 6):
        c = ws.cell(row=row_idx, column=col)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = BORDER
    row_idx += 1
    n += 1

ws.column_dimensions["A"].width = 5
ws.column_dimensions["B"].width = 20
ws.column_dimensions["C"].width = 42
ws.column_dimensions["D"].width = 55
ws.column_dimensions["E"].width = 38

ws.auto_filter.ref = f"A1:E{row_idx-1}"

# ---------- Sheet 2: How To Use (n8n notes) ----------
ws2 = wb.create_sheet("How To Use (n8n)")
notes = [
    ("GW Fins Chatbot — FAQ Knowledge Base", True, 14),
    ("", False, 10),
    ("Purpose:", True, 11),
    ("This sheet is a starter knowledge base of the most common questions restaurant guests ask a chatbot. Use it as your RAG/knowledge-base source or as intent-training data for the n8n chatbot flow.", False, 10),
    ("", False, 10),
    ("How to use in n8n:", True, 11),
    ("1. Export this sheet as CSV or connect the Google Sheet directly via the Google Sheets node.", False, 10),
    ("2. Feed the 'Customer Question (Intent)' + 'Suggested Chatbot Answer' columns into your vector store / RAG node (e.g. Pinecone, Supabase, or a Google Sheets lookup + AI Agent node) as document chunks.", False, 10),
    ("3. Use the 'Category' column to route intents (e.g. Reservations -> booking flow, Menu -> menu lookup, Private Events -> lead capture to staff).", False, 10),
    ("4. Rows highlighted in yellow under 'Notes / Needs Client Confirmation' contain placeholder or best-guess answers — confirm exact policy with GW Fins management before publishing live.", False, 10),
    ("5. For live data (today's menu, current availability), wire the chatbot to the real menu page (gwfins.com/menu) or reservation system API rather than hardcoding answers that go stale.", False, 10),
    ("", False, 10),
    ("Key facts already verified from gwfins.com / Google Business listing:", True, 11),
    ("Address: 808 Bienville St, New Orleans, LA 70112, United States", False, 10),
    ("Phone: (504) 581-3467", False, 10),
    ("Hours: Dinner nightly from 5:00 PM (menu finalized after 4:00 PM daily catch arrival)", False, 10),
    ("Located: DH Holmes building, 4 blocks from Saenger Theatre, edge of French Quarter/CBD", False, 10),
    ("Price range: $40-$100 per person", False, 10),
    ("Rating: 4.8 stars (5,507+ Google reviews)", False, 10),
    ("Anniversary: Celebrating 25 years (2001-2026)", False, 10),
]
r = 1
for text, bold, size in notes:
    c = ws2.cell(row=r, column=1, value=text)
    c.font = Font(name=FONT_NAME, bold=bold, size=size, color="1F3864" if bold else "000000")
    c.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
ws2.column_dimensions["A"].width = 110
for i in range(1, r):
    ws2.row_dimensions[i].height = 16 if notes[i-1][1] else 28

wb.save(r"C:\Users\user\Downloads\GW Fins restaurant\chatbot\GW_Fins_Chatbot_FAQ.xlsx")
print("saved", row_idx-2, "FAQ rows")
